#!/usr/bin/env python3
"""Populate the hidden CSR Store Name column on NVS List (NEW) and NSO List (NEW).

Maps Store # -> CSR Store Name from a combined NVS+NSO spreadsheet, then writes
to the existing hidden TEXT_NUMBER column "CSR Store Name" on each sheet via the
Smartsheet REST API. Only cells whose row Store # is in the mapping are touched;
unmatched, blank, and TBD rows are left untouched. Idempotent: if a row's
existing CSR Store Name already equals the mapped value, the row is skipped.

Auth:
  Reads $SMARTSHEET_ACCESS_TOKEN. If unset, falls back to AWS Secrets Manager
  (secret id from $SMARTSHEET_AWS_SECRET_ID, default "smartsheet/access_token",
  region from $AWS_REGION, default us-west-2).

CLI:
  --xlsx        path to the source workbook (defaults to data/Store_Names_Updated.xlsx
                under --working-dir, then falls back to "Store Names Updated.xlsx"
                in the working-dir if the underscore-named file is absent)
  --working-dir base directory used to resolve relative paths
  --dry-run     compute and print everything, but send no PUT requests
"""

from __future__ import annotations

import argparse
import json
import os
import random
import sys
import time
from pathlib import Path
from typing import Any

import openpyxl
import requests

API_BASE = "https://api.smartsheet.com/2.0"

SHEETS = [
    {
        "label": "NVS List (NEW)",
        "sheet_id": 8700079041892228,
        "csr_column_id": 6363478419083140,
    },
    {
        "label": "NSO List (NEW)",
        "sheet_id": 328848479571844,
        "csr_column_id": 6957589299761028,
    },
]

CSR_COLUMN_TITLE = "CSR Store Name"
STORE_NUMBER_TITLES = {"Store #", "Store#", "Store Number", "Store No"}
BLANK_STORE_TOKENS = {"", "TBD", "N/A", "NA", "NONE", "-", "TBA"}
PUT_CHUNK_SIZE = 400


def get_token() -> str:
    token = os.environ.get("SMARTSHEET_ACCESS_TOKEN")
    if token:
        return token.strip()

    secret_id = os.environ.get("SMARTSHEET_AWS_SECRET_ID", "smartsheet/access_token")
    region = os.environ.get("AWS_REGION", "us-west-2")

    try:
        import boto3  # type: ignore
    except ImportError:
        sys.exit(
            "SMARTSHEET_ACCESS_TOKEN not set and boto3 is not installed for the "
            "Secrets Manager fallback. pip install boto3 or export the token."
        )

    client = boto3.client("secretsmanager", region_name=region)
    resp = client.get_secret_value(SecretId=secret_id)
    raw = resp.get("SecretString") or ""

    # Secret may be a raw token or a JSON blob like {"token": "..."} or
    # {"SMARTSHEET_ACCESS_TOKEN": "..."} or {"api_token": "..."}.
    raw_stripped = raw.strip()
    if raw_stripped.startswith("{"):
        try:
            obj = json.loads(raw_stripped)
        except json.JSONDecodeError:
            return raw_stripped
        for key in ("token", "SMARTSHEET_ACCESS_TOKEN", "smartsheet_access_token",
                    "api_token", "access_token", "value"):
            if key in obj and obj[key]:
                return str(obj[key]).strip()
        # Fall through: if there's exactly one string value, use it.
        str_values = [v for v in obj.values() if isinstance(v, str) and v]
        if len(str_values) == 1:
            return str_values[0].strip()
        sys.exit(
            f"Could not find a token key in secret {secret_id}; "
            f"keys present: {sorted(obj.keys())}"
        )
    return raw_stripped


def normalize_store_number(raw: Any) -> str | None:
    """Coerce a raw cell value into a canonical store-number string.

    Returns None for blanks/TBD/etc. so the caller can skip those rows.
    """
    if raw is None:
        return None
    if isinstance(raw, float):
        # openpyxl/Smartsheet often return ints as floats; drop the trailing .0.
        if raw != raw:  # NaN
            return None
        if raw.is_integer():
            raw = int(raw)
    s = str(raw).strip()
    if not s:
        return None
    upper = s.upper()
    if upper in BLANK_STORE_TOKENS:
        return None
    # Strip a trailing ".0" leftover from text-typed floats.
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    # Drop leading zeros but keep a single zero if the value really is "0".
    if s.isdigit():
        s = s.lstrip("0") or "0"
    return s


def load_mapping(xlsx_path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    mapping: dict[str, str] = {}
    conflicts: list[tuple[str, str, str, str]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        # Header is "Store #, Store Name" per spec — fail loudly if not.
        cols = [(c or "").strip() if isinstance(c, str) else c for c in header]
        try:
            store_idx = next(i for i, c in enumerate(cols) if c in STORE_NUMBER_TITLES)
            name_idx = next(i for i, c in enumerate(cols)
                            if isinstance(c, str) and "name" in c.lower())
        except StopIteration:
            sys.exit(f"Sheet '{sheet_name}' in {xlsx_path} is missing expected headers; "
                     f"got {cols!r}")
        for row in rows:
            if not row or all(v is None for v in row):
                continue
            store = normalize_store_number(row[store_idx])
            if store is None:
                continue
            name_raw = row[name_idx]
            if name_raw is None:
                continue
            name = str(name_raw).strip()
            if not name:
                continue
            if store in mapping and mapping[store] != name:
                conflicts.append((store, mapping[store], name, sheet_name))
                # Last-write-wins, but report it.
            mapping[store] = name
    if conflicts:
        print(f"[warn] {len(conflicts)} store # conflicts between sheets; "
              f"last-write-wins. Examples:")
        for s, a, b, where in conflicts[:5]:
            print(f"  {s}: '{a}' -> '{b}' (from {where})")
    return mapping


class SmartsheetClient:
    def __init__(self, token: str):
        self.session = requests.Session()
        self.session.headers.update({
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        })

    def _request(self, method: str, path: str, **kwargs) -> dict:
        url = f"{API_BASE}{path}"
        attempt = 0
        while True:
            attempt += 1
            try:
                resp = self.session.request(method, url, timeout=60, **kwargs)
            except requests.RequestException as e:
                if attempt >= 6:
                    raise
                wait = min(2 ** attempt, 30) + random.uniform(0, 1)
                print(f"[retry] {method} {path}: {e!r}; sleeping {wait:.1f}s")
                time.sleep(wait)
                continue

            if resp.status_code in (429,) or 500 <= resp.status_code < 600:
                if attempt >= 6:
                    resp.raise_for_status()
                wait = min(2 ** attempt, 30) + random.uniform(0, 1)
                print(f"[retry] {method} {path}: HTTP {resp.status_code}; "
                      f"sleeping {wait:.1f}s")
                time.sleep(wait)
                continue

            if not resp.ok:
                # Surface server error body to help debug 4xx.
                raise requests.HTTPError(
                    f"{method} {path} -> {resp.status_code}: {resp.text[:500]}",
                    response=resp,
                )
            if not resp.content:
                return {}
            return resp.json()

    def get_sheet(self, sheet_id: int) -> dict:
        # Page through if the sheet is huge. NVS/NSO are well under 5k rows,
        # but defend against future growth.
        all_rows: list[dict] = []
        columns: list[dict] | None = None
        name: str | None = None
        page = 1
        page_size = 5000
        while True:
            data = self._request(
                "GET", f"/sheets/{sheet_id}",
                params={"pageSize": page_size, "page": page,
                        "include": "objectValue"},
            )
            if columns is None:
                columns = data.get("columns", [])
                name = data.get("name")
            all_rows.extend(data.get("rows", []))
            total_pages = data.get("totalPages", 1)
            if page >= total_pages:
                break
            page += 1
        return {"name": name, "columns": columns or [], "rows": all_rows}

    def update_rows(self, sheet_id: int, updates: list[dict]) -> int:
        """PUT row updates in chunks of PUT_CHUNK_SIZE; return rows updated."""
        total = 0
        for i in range(0, len(updates), PUT_CHUNK_SIZE):
            chunk = updates[i:i + PUT_CHUNK_SIZE]
            self._request("PUT", f"/sheets/{sheet_id}/rows", data=json.dumps(chunk))
            total += len(chunk)
        return total


def find_store_column_id(columns: list[dict]) -> int:
    for c in columns:
        if c.get("title") in STORE_NUMBER_TITLES:
            return c["id"]
    raise SystemExit(
        "Could not find a 'Store #' column on the sheet; "
        f"columns present: {[c.get('title') for c in columns]}"
    )


def find_or_warn_csr_column(columns: list[dict], expected_id: int) -> int:
    for c in columns:
        if c.get("id") == expected_id:
            if c.get("title") != CSR_COLUMN_TITLE:
                print(f"[warn] column id {expected_id} title is "
                      f"{c.get('title')!r}, not {CSR_COLUMN_TITLE!r}")
            return c["id"]
    # The pre-created column has gone missing — fall back to title lookup,
    # but loudly. The spec says we should NOT add another one.
    for c in columns:
        if c.get("title") == CSR_COLUMN_TITLE:
            print(f"[warn] expected CSR column id {expected_id} not found; "
                  f"matched by title to id {c['id']}")
            return c["id"]
    raise SystemExit(
        f"Neither column id {expected_id} nor a column titled "
        f"{CSR_COLUMN_TITLE!r} found. Refusing to add a new column."
    )


def process_sheet(client: SmartsheetClient, cfg: dict, mapping: dict[str, str],
                  dry_run: bool) -> dict:
    label = cfg["label"]
    sheet_id = cfg["sheet_id"]
    expected_csr_id = cfg["csr_column_id"]

    print(f"\n=== {label} (id {sheet_id}) ===")
    sheet = client.get_sheet(sheet_id)
    columns = sheet["columns"]
    rows = sheet["rows"]
    print(f"  fetched {len(rows)} rows, {len(columns)} columns "
          f"(sheet name: {sheet['name']!r})")

    csr_col_id = find_or_warn_csr_column(columns, expected_csr_id)
    store_col_id = find_store_column_id(columns)

    matched = 0
    already_correct = 0
    unmatched: list[str] = []
    blank = 0
    changes: list[dict] = []

    for row in rows:
        store_cell = next((c for c in row.get("cells", [])
                           if c.get("columnId") == store_col_id), None)
        csr_cell = next((c for c in row.get("cells", [])
                         if c.get("columnId") == csr_col_id), None)

        if store_cell is None:
            blank += 1
            continue
        store_raw = store_cell.get("value")
        if store_raw is None:
            # Some rows use displayValue but no value when type drifts.
            store_raw = store_cell.get("displayValue")
        store = normalize_store_number(store_raw)
        if store is None:
            blank += 1
            continue
        if store not in mapping:
            unmatched.append(store)
            continue

        desired = mapping[store]
        current = (csr_cell or {}).get("value")
        if current is not None and str(current).strip() == desired:
            already_correct += 1
            matched += 1
            continue

        changes.append({
            "id": row["id"],
            "cells": [{"columnId": csr_col_id, "value": desired}],
        })
        matched += 1

    print(f"  matched:        {matched}")
    print(f"  already correct (skipped): {already_correct}")
    print(f"  to update:      {len(changes)}")
    print(f"  unmatched:      {len(unmatched)}")
    print(f"  blank/TBD:      {blank}")
    if unmatched:
        sample = sorted(set(unmatched))[:15]
        print(f"  unmatched sample: {sample}")

    if not changes:
        print("  -> no changes to send")
        return {"matched": matched, "updated": 0, "unmatched": len(unmatched)}

    if dry_run:
        print("  -> DRY RUN: not sending PUT")
        # Show first few intended writes for transparency.
        for change in changes[:3]:
            print(f"     would set row {change['id']} CSR -> "
                  f"{change['cells'][0]['value']!r}")
        return {"matched": matched, "updated": 0,
                "would_update": len(changes), "unmatched": len(unmatched)}

    updated = client.update_rows(sheet_id, changes)
    print(f"  -> updated {updated} rows")
    return {"matched": matched, "updated": updated, "unmatched": len(unmatched)}


def parse_args() -> argparse.Namespace:
    default_wd = r"C:\Users\vanbu\OneDrive\Desktop\Nike Blueprint Support"
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--xlsx", default=None,
                   help="Path to the Store_Names workbook. If omitted, looks for "
                        "data/Store_Names_Updated.xlsx (then "
                        "'Store Names Updated.xlsx') under --working-dir.")
    p.add_argument("--working-dir", default=default_wd)
    p.add_argument("--dry-run", action="store_true")
    return p.parse_args()


def resolve_xlsx(args: argparse.Namespace) -> Path:
    if args.xlsx:
        p = Path(args.xlsx)
        if not p.exists():
            sys.exit(f"--xlsx path does not exist: {p}")
        return p
    wd = Path(args.working_dir)
    candidates = [
        wd / "data" / "Store_Names_Updated.xlsx",
        wd / "Store_Names_Updated.xlsx",
        wd / "Store Names Updated.xlsx",
        # Fallback: this script's own repo data dir, if run from elsewhere.
        Path(__file__).resolve().parent.parent / "data" / "Store_Names_Updated.xlsx",
    ]
    for c in candidates:
        if c.exists():
            return c
    sys.exit("Could not find Store_Names_Updated.xlsx. Tried: "
             + ", ".join(str(c) for c in candidates))


def main() -> int:
    args = parse_args()
    xlsx_path = resolve_xlsx(args)
    print(f"loading mapping from {xlsx_path}")
    mapping = load_mapping(xlsx_path)
    print(f"  {len(mapping)} unique Store # -> CSR Store Name entries")

    token = get_token()
    client = SmartsheetClient(token)

    results = []
    for cfg in SHEETS:
        results.append((cfg["label"], process_sheet(client, cfg, mapping, args.dry_run)))

    print("\n=== summary ===")
    for label, r in results:
        print(f"  {label}: {r}")
    print(f"  mode: {'DRY RUN' if args.dry_run else 'LIVE WRITE'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
